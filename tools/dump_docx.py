import sys
from docx import Document
path = r"D:\Projects\Nikola Ispiti\Kteh\Petar_Petrović_2022_0102_Domaći_2.docx"
doc = Document(path)
print("=== PARAGRAPHS (" + str(len(doc.paragraphs)) + ") ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        print("[" + str(i) + "] (" + p.style.name + ") " + t[:140])
print("\n=== TABLES (" + str(len(doc.tables)) + ") ===")
for ti, tbl in enumerate(doc.tables):
    print("-- table " + str(ti) + ": " + str(len(tbl.rows)) + "r x " + str(len(tbl.columns)) + "c --")
    for ri, row in enumerate(tbl.rows):
        cells = [c.text.strip()[:50] for c in row.cells]
        print("  r" + str(ri) + ": " + str(cells))
