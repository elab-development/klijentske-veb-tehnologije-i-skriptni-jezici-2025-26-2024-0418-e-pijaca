# -*- coding: utf-8 -*-
"""Fill the FON Domaći 2 DOCX template for the E-Pijaca project."""
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

TEMPLATE = r"D:\Projects\Nikola Ispiti\Kteh\Petar_Petrović_2022_0102_Domaći_2.docx"
OUTPUT   = r"D:\Projects\Nikola Ispiti\Kteh\Nikola_Markovic_2024_0560_Domaci_2.docx"
SHOTS    = r"C:\Users\MarkoMarkovic\AppData\Local\hermes\cache\screenshots"
REPO     = "https://github.com/elab-development/klijentske-veb-tehnologije-2024-2024-0560-e-pijaca"

# (filename, title, description)
SCREENS = [
    ("browser_screenshot_0001f8032d7f4fad9a81acd566e7a7af.png",
     "Početna stranica",
     "Prikazuje hero sekciju sa pozivom na akciju („Pogledaj ponudu”, „Postani član”), mrežu kategorija "
     "(Voće, Povrće, Mlečni, Med, Jaja, Začini), preporučene proizvode sa ocenama i cenama, te sekciju "
     "sezonskih akcija. Dugme za promenu valute (RSD ⇄ EUR) dostupno je u sekciji preporuka."),
    ("browser_screenshot_7ebb5b775ef04818a19f74a308e929aa.png",
     "Spisak proizvoda (Marketplace)",
     "Prikazuje sve proizvode (12) sa bočnim panelom za filtere (kategorija, cena, region, sertifikat), "
     "sortiranjem, pretragom i paginacijom (6 proizvoda po stranici, 2 strane). Svaka kartica prikazuje "
     "oznaku kategorije, naziv, proizvođača, ocenu i cenu, sa dugmetom „Dodaj u korpu”."),
    ("browser_screenshot_65d09a66a4ef4a1c80fbefdfcabdc3f1.png",
     "Pojedinačni proizvod",
     "Prikazuje galeriju slika, oznaku kategorije, naziv proizvoda, ocenu sa brojem recenzija i brojem "
     "prodatih komada, cenu sa precrtanom starom cenom i procentom popusta, birač količine (Smanji/Povećaj), "
     "dugme „Dodaj u korpu” sa ukupnom cenom, te tabove: Opis, Detalji i Recenzije."),
    ("browser_screenshot_696122ea59364fd7a80d3fbf4bd21d25.png",
     "Korpa i naručivanje",
     "Prikazuje četvorokoračni proces (Korpa → Dostava → Plaćanje → Potvrda). Proizvodi su grupisani po "
     "proizvođaču (Voćnjak Mihajlović, Srpsko selo dairy, Malinarski raj), sa biračem količine i ukidanjem. "
     "Bočna rekapitulacija prikazuje vrednost proizvoda, PDV (20%), dostavu, ukupno i polje za kod za popust "
     "(PIJACA10 = 10%, SEZONA15 = 15%)."),
    ("browser_screenshot_59e8db81f0f846978cc4ee172734c38e.png",
     "Prijava",
     "Forma za prijavu sa poljima za email i lozinku (sa ikonama), dugmetom „Prijavi se” i linkom za "
     "registraciju. Demo režim: bilo koji email i lozinka radi. Nakon prijave, korisnik ostaje prijavljen "
     "zahvaljujući trajnoj sesiji u localStorage."),
    ("browser_screenshot_fdec3161dd5b44c9978f4d68b3395be4.png",
     "Registracija",
     "Forma za registraciju sa poljima: ime, prezime, email adresa, telefon, lozinka i adresa dostave, "
     "sa dugmetom „Registruj se” i linkom za prijavu. Nakon registracije, nalog se čuva u localStorage i "
     "korisnik biva prijavljen."),
    ("browser_screenshot_45b5b439c5694f3e969a87b71fd8dbab.png",
     "Profil korisnika",
     "Prikazuje avatar sa inicijalima, ime i email korisnika, karticu „Podaci o nalogu” (telefon, adresa, "
     "član od, broj proizvoda u korpi) i karticu „Istorija porudžbina” sa listom porudžbina i statusima "
     "(Dostavljeno, Poslato, U obradi). Dugme „Odjavi se” vrši odjavu."),
]

def insert_after(ref, text="", mono=False, size=11, bold=False):
    new_p = ref._p.makeelement(qn('w:p'), {})
    ref._p.addnext(new_p)
    np = Paragraph(new_p, ref._parent)
    if text:
        run = np.add_run(text)
        if mono:
            run.font.name = 'Consolas'
            run.font.size = Pt(size)
        if bold:
            run.bold = True
    return np

def insert_image(ref, path, width=5.7):
    np = insert_after(ref)
    np.add_run().add_picture(path, width=Inches(width))
    return np

# ---- real code snippets (verified against source) ----
SNIP_CART = (
    "// models/Cart.ts — klasa korpe sa poslovnim logikama\n"
    "export class Cart {\n"
    "  items: ICartItem[];\n"
    "  discountCode: string | null;\n"
    "  discountRate: number;\n"
    "  constructor(items: ICartItem[] = [], discountCode: string | null = null) {\n"
    "    this.items = items;\n"
    "    this.discountCode = discountCode;\n"
    "    this.discountRate = discountCode ? (DISCOUNT_CODES[discountCode] ?? 0) : 0;\n"
    "  }\n"
    "  addItem(productId: number, quantity = 1): void {\n"
    "    const existing = this.items.find((i) => i.productId === productId);\n"
    "    if (existing) existing.quantity += quantity;\n"
    "    else this.items.push({ productId, quantity });\n"
    "  }\n"
    "  getSubtotal(priceOf: (id: number) => number): number {\n"
    "    return this.items.reduce((s, i) => s + priceOf(i.productId) * i.quantity, 0);\n"
    "  }\n"
    "  getDiscountedTotal(priceOf): number {\n"
    "    return Math.round(this.getSubtotal(priceOf) * (1 - this.discountRate));\n"
    "  }\n"
    "}"
)
SNIP_CTX = (
    "// context/CartContext.tsx — React kontekst za stanje korpe\n"
    "export function CartProvider({ children }) {\n"
    "  const [items, setItems] = useLocalStorage<ICartItem[]>('epijaca-cart-items', []);\n"
    "  const [discountCode, setDiscountCode] = useLocalStorage<string | null>('epijaca-cart-code', null);\n"
    "  const cart = useMemo(() => new Cart(items, discountCode), [items, discountCode]);\n"
    "  const priceOf = (id) => effectivePrice(getProductById(id));\n"
    "  // addItem: radi na KOPIJI niza da React detektuje promenu stanja\n"
    "  const addItem = (productId, qty = 1) => {\n"
    "    const c = new Cart([...items], discountCode);\n"
    "    c.addItem(productId, qty);\n"
    "    setItems(c.items);\n"
    "  };\n"
    "  // ...removeItem, updateQty, applyDiscount, clear\n"
    "  return <CartContext.Provider value={{ items, count, subtotal, total, addItem, ... }}>{children}</CartContext.Provider>;\n"
    "}"
)

def main():
    doc = Document(TEMPLATE)

    # 1) student table
    tbl = doc.tables[0]
    students = [
        ("Nikola", "Marković", "2024/0560"),
        ("Jovan",  "Luković",  "2024/0418"),
        ("Iva",    "Krstev",   "2024/0481"),
    ]
    for i, (ime, prez, idx) in enumerate(students, start=1):
        tbl.rows[i].cells[0].text = ime
        tbl.rows[i].cells[1].text = prez
        tbl.rows[i].cells[2].text = idx
    tbl.rows[4].cells[1].text = REPO

    # capture paragraph refs BEFORE inserting (indices will shift)
    p8  = doc.paragraphs[8]
    p10 = doc.paragraphs[10]
    p12 = doc.paragraphs[12]
    p13 = doc.paragraphs[13]

    # 2) Tema
    p8.text = (
        "E-Pijaca je tematska web aplikacija — elektronska pijaca koja povezuje lokalne srpske proizvođače "
        "(voćnjake, mlekare, pčelare, rasadnike) sa kupcima. Korisnici mogu pregledati katalog domaćih "
        "proizvoda (voće, povrće, mlečni proizvodi, med, jaja, začini), filtrisati ih po kategoriji, ceni, "
        "regionu i sertifikatu, videti detalje svakog proizvoda sa ocenama i recenzijama, dodavati ih u "
        "korpu, primeniti kod za popust i završiti četvorokoračni proces naručivanja (korpa → dostava → "
        "plaćanje → potvrda). Aplikacija nudi prijavu i registraciju sa trajnom sesijom, korisnički profil "
        "sa istorijom porudžbina, i konverziju cena RSD ⇄ EUR preko spoljnog API-ja za kursnu listu. "
        "Dizajn je veran Figma prototipu (zelena #2D7A3E, krem pozadina #FAF8F3) i potpuno je responsivan."
    )

    # 3) Implementacija (intro on p10, then insert details + code)
    p10.text = (
        "Aplikacija je razvijena u Vite + React 18 + TypeScript tehnologiji, sa react-router-dom v6 za "
        "ruting i čistim CSS-om (dizajn tokeni iz Fige) za stilizaciju. Arhitektura je slojevita: models "
        "(TypeScript klase), data (repository pattern), services (2 spoljna REST API-ja), context (React "
        "state), hooks, components i pages."
    )
    ref = p10
    blocks = [
        ("p", "Domen modelovan je klasama Product, Cart, User, Order i interfejsima IProduct, ICartItem, "
              "IProducer, IReview, IOrder, IUser, IFilterOptions, IRepository<T>. Primer klase Cart:"),
        ("c", SNIP_CART),
        ("p", "Custom hook useLocalStorage persistuje stanje u browser memoriji (korpa, korisnik, valuta, "
              "popust); useDebounce odlaže pretragu; useFetch radi fetch sa statusima loading/error. "
              "React konteksti upravljaju globalnim stanjem: CartContext (korpa, popusti, PDV, dostava), "
              "AuthContext (prijava/sesija/rekonstrukcija User instance), CurrencyContext (RSD/EUR). "
              "Primer CartContext.addItem:"),
        ("c", SNIP_CTX),
        ("p", "Ruting (App.tsx) koristi react-router-dom v6 sa 7 ruta: / (Početna), /proizvodi (Spisak), "
              "/proizvod/:id (Detalj), /prijava, /registracija, /profil (zaštićena ruta), /korpa. "
              "Spoljni API-ji: productService integracija sa Fake Store API (https://fakestoreapi.com) i "
              "currencyService sa open.er-api.com za kursnu listu. Paginacija, filteri i sortiranje se "
              "sinhronizuju sa URL query parametrima (useSearchParams)."),
    ]
    for kind, text in blocks:
        ref = insert_after(ref, text, mono=(kind == "c"), size=8 if kind == "c" else 11)

    # 4) Korisničko uputstvo (intro on p12/p13, then 7 screenshots)
    p12.text = ("U nastavku su prikazani screenshotovi svih stranica aplikacije u pretraživaču, uz opis "
                "funkcionalnosti koje se na svakoj stranici nude korisniku.")
    p13.text = ""
    ref = p13
    for i, (fname, title, desc) in enumerate(SCREENS, 1):
        t = insert_after(ref, f"Slika {i}. {title}", bold=True)
        imgp = insert_image(t, os.path.join(SHOTS, fname), width=5.7)
        ref = insert_after(imgp, desc)

    doc.save(OUTPUT)
    sz = os.path.getsize(OUTPUT)
    print(f"DOCX saved: {OUTPUT}")
    print(f"Size: {sz:,} bytes")
    print(f"Screenshots embedded: {len(SCREENS)}")
    print(f"Students in table: {len(students)}")

if __name__ == "__main__":
    main()
